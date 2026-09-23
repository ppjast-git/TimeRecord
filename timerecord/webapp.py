"""FastAPI dashboard: przegląd czasu pracy w przeglądarce.

Endpointy:
  GET /                  -> strona HTML dashboardu
  GET /api/today         -> JSON: suma dzisiaj + per-app
  GET /api/day/{date}    -> JSON: suma dla dnia (YYYY-MM-DD)
  GET /api/week          -> JSON: ostatnie 7 dni
  GET /api/events?date=  -> JSON: ostatnie zdarzenia danego dnia
  GET /api/now           -> JSON: aktualna próbka z collectora
  POST /api/pause        -> pauza collectora
  POST /api/resume       -> wznowienie collectora
"""
from __future__ import annotations

import datetime as _dt
import io
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import HTMLResponse, JSONResponse, Response

from .collector import Collector
from .config import SETTINGS
from .storage import Storage

_TEMPLATES_DIR = Path(__file__).parent.parent / "templates"
_STATIC_DIR = Path(__file__).parent.parent / "static"


def _fmt_duration(seconds: float) -> str:
    """Sekundy -> '1h 23m' / '12m 30s' / '45s'."""
    s = int(round(seconds))
    h, s = divmod(s, 3600)
    m, s = divmod(s, 60)
    if h:
        return f"{h}h {m:02d}m"
    if m:
        return f"{m}m {s:02d}s"
    return f"{s}s"


def create_app(storage: Storage, collector: Collector) -> FastAPI:
    app = FastAPI(title="TimeRecord", docs_url=None, redoc_url=None)

    @app.get("/", response_class=HTMLResponse)
    async def index() -> HTMLResponse:
        html = (_TEMPLATES_DIR / "dashboard.html").read_text(encoding="utf-8")
        return HTMLResponse(html)

    @app.get("/api/today")
    async def api_today() -> JSONResponse:
        return _enrich(storage.daily_summary())

    @app.get("/api/day/{date}")
    async def api_day(date: str) -> JSONResponse:
        try:
            d = _dt.datetime.fromisoformat(date).astimezone()
        except ValueError:
            raise HTTPException(400, "Niepoprawna data (użyj YYYY-MM-DD)")
        return _enrich(storage.daily_summary(d))

    @app.get("/api/week")
    async def api_week() -> JSONResponse:
        days = storage.week_summary()
        return JSONResponse({"days": [_enrich(d) for d in days]})

    @app.get("/api/events")
    async def api_events(date: Optional[str] = None, limit: int = 50) -> JSONResponse:
        if date:
            try:
                d = _dt.datetime.fromisoformat(date).astimezone()
            except ValueError:
                raise HTTPException(400, "Niepoprawna data")
        else:
            d = None
        events = storage.recent_events(d, limit=limit)
        for e in events:
            e["dur_human"] = _fmt_duration(e.get("dur_seconds") or 0)
        return JSONResponse({"events": events})

    @app.get("/api/day/{date}/app/{app}/top-titles")
    async def api_top_titles(
        date: str,
        app: str,
        limit: int = 8,
    ) -> JSONResponse:
        try:
            d = _dt.datetime.fromisoformat(date).astimezone()
        except ValueError:
            raise HTTPException(400, "Niepoprawna data (użyj YYYY-MM-DD)")
        limit = max(1, min(int(limit), 20))
        rows = storage.top_titles_for_day_and_app(day=d, app=app, limit=limit)
        return JSONResponse({"app": app, "day": d.date().isoformat(), "items": rows})

    @app.get("/api/day/{date}/app/{app}/entities")
    async def api_app_entities(
        date: str,
        app: str,
        limit: int = 100,
    ) -> JSONResponse:
        try:
            d = _dt.datetime.fromisoformat(date).astimezone()
        except ValueError:
            raise HTTPException(400, "Niepoprawna data (użyj YYYY-MM-DD)")
        limit = max(1, min(int(limit), 200))
        entities = storage.entities_for_day(day=d, app=app, limit=limit)
        return JSONResponse({"app": app, "day": d.date().isoformat(), "entities": entities})

    @app.get("/api/day/{date}/app/{app}/levels")
    async def api_app_levels(
        date: str,
        app: str,
        top: int = 50,
    ) -> JSONResponse:
        # Alias dla kompatybilności wstecznej
        return await api_app_entities(date=date, app=app, limit=top)

    @app.get("/api/day/{date}/projects")
    async def api_day_projects(date: str) -> JSONResponse:
        try:
            d = _dt.datetime.fromisoformat(date).astimezone()
        except ValueError:
            raise HTTPException(400, "Niepoprawna data (użyj YYYY-MM-DD)")
        projects = storage.projects_for_day(day=d)
        return JSONResponse({"day": d.date().isoformat(), "projects": projects})

    @app.get("/api/month/{year}/{month}")
    async def api_month(year: int, month: int) -> JSONResponse:
        if not (1 <= month <= 12) or not (1900 <= year <= 2100):
            raise HTTPException(400, "Niepoprawny rok/miesiąc")
        days = storage.month_summary(year, month)
        return JSONResponse({
            "year": year,
            "month": month,
            "days": [
                {"day": s["day"],
                 "total_seconds": s["total_seconds"],
                 "idle_seconds": s["idle_seconds"]}
                for s in days
            ],
        })

    @app.get("/api/day/{date}/export.xlsx")
    async def api_export_day(date: str) -> Response:
        try:
            d = _dt.datetime.fromisoformat(date).astimezone()
        except ValueError:
            raise HTTPException(400, "Niepoprawna data (użyj YYYY-MM-DD)")
        try:
            from openpyxl import Workbook
        except ImportError:
            raise HTTPException(500, "Brak zależności openpyxl (pip install openpyxl)")

        summary = storage.daily_summary(d)
        events = storage.recent_events(d, limit=10000)
        total = summary.get("total_seconds") or 0
        idle_s = summary.get("idle_seconds") or 0

        def _secs(v) -> int:
            """Czas w sekundach zaokrąglony do pełnej wartości całkowitej."""
            return int(round(v or 0))

        def _local(s: Optional[str]) -> Optional[str]:
            """ISO (UTC) -> 'YYYY-MM-DD HH:MM:SS' w czasie lokalnym."""
            if not s:
                return s
            try:
                return _dt.datetime.fromisoformat(s).astimezone().strftime("%Y-%m-%d %H:%M:%S")
            except ValueError:
                return s

        wb = Workbook()

        ws = wb.active
        ws.title = "Aplikacje"
        ws.append(["Aplikacja", "Czas (s)", "Czas", "% aktywnego", "Zdarzeń"])
        for a in summary.get("apps", []):
            dur = a.get("dur_seconds") or 0
            ws.append([
                a["app"],
                _secs(dur),
                _fmt_duration(dur),
                round(dur / total * 100, 1) if total else 0,
                a.get("n_events", 0),
            ])
        ws.append([])
        ws.append(["SUMA aktywnego", _secs(total), _fmt_duration(total), 100.0 if total else 0, ""])
        ws.append(["Idle (AFK)", _secs(idle_s), _fmt_duration(idle_s), "", ""])
        for col, w in {"A": 24, "B": 12, "C": 12, "D": 14, "E": 10}.items():
            ws.column_dimensions[col].width = w

        ws2 = wb.create_sheet("Zdarzenia")
        ws2.append(["Start", "Koniec", "Aplikacja", "Tytuł okna", "Zakładka",
                    "Przeglądarka", "Idle", "Czas (s)", "Czas"])
        for e in events:
            dur = e.get("dur_seconds") or 0
            ws2.append([
                _local(e.get("ts_start")), _local(e.get("ts_end")), e.get("app"),
                e.get("title"), e.get("tab"), e.get("browser"),
                "tak" if e.get("idle") else "",
                _secs(dur), _fmt_duration(dur),
            ])
        for col, w in {"A": 20, "B": 20, "C": 16, "D": 60, "E": 44,
                       "F": 16, "G": 6, "H": 10, "I": 10}.items():
            ws2.column_dimensions[col].width = w

        ws3 = wb.create_sheet("Obiekty i zadania")
        ws3.append(["Aplikacja", "Obiekt / Zadanie", "Kategoria", "Projekt", "Czas (s)",
                    "Czas", "% aplikacji", "Wariantów okien", "Zdarzeń"])
        app_totals = {a["app"]: a["dur_seconds"] for a in summary.get("apps", [])}
        for ent in storage.entities_for_day(d, limit=2000):
            dur = ent.get("dur_seconds") or 0
            app_t = app_totals.get(ent["app"], 0)
            pct = round(dur / app_t * 100, 1) if app_t else 0
            ws3.append([
                ent["app"], ent["entity"], ent["category"], ent["project"] or "",
                _secs(dur), _fmt_duration(dur), pct,
                ent.get("n_variants", 0), ent.get("n_events", 0),
            ])
        for col, w in {"A": 16, "B": 44, "C": 22, "D": 16, "E": 10, "F": 10,
                       "G": 12, "H": 16, "I": 10}.items():
            ws3.column_dimensions[col].width = w

        ws4 = wb.create_sheet("Projekty (zbiorczo)")
        ws4.append(["Projekt", "Czas (s)", "Czas", "% aktywnego dnia", "Aplikacje", "Zdarzeń", "Główne zadania"])
        for pr in storage.projects_for_day(d):
            dur = pr.get("dur_seconds") or 0
            pct = round(dur / total * 100, 1) if total else 0
            top_desc = ", ".join(f"{x['entity']} ({_fmt_duration(x['dur_seconds'])})" for x in pr.get("top_entities", []))
            ws4.append([
                pr["project"], _secs(dur), _fmt_duration(dur), pct,
                ", ".join(pr.get("apps", [])), pr.get("n_events", 0), top_desc,
            ])
        for col, w in {"A": 18, "B": 10, "C": 10, "D": 16, "E": 28, "F": 10, "G": 60}.items():
            ws4.column_dimensions[col].width = w

        ws5 = wb.create_sheet("Wszystkie tytuły (szczegóły)")
        ws5.append(["Aplikacja", "Tytuł / zakładka", "Czas (s)", "Czas", "Zdarzeń"])
        for it in storage.all_titles_for_day(d):
            dur = it.get("dur_seconds") or 0
            ws5.append([it["app"], it.get("key"), _secs(dur), _fmt_duration(dur),
                        it.get("n_events", 0)])
        for col, w in {"A": 16, "B": 90, "C": 10, "D": 10, "E": 8}.items():
            ws5.column_dimensions[col].width = w

        buf = io.BytesIO()
        wb.save(buf)
        filename = f"timerecord-{d.date().isoformat()}.xlsx"
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    @app.get("/api/now")
    async def api_now() -> JSONResponse:
        s = collector.last_sample
        return JSONResponse({
            "paused": collector.is_paused,
            "sample": s,
            "host": SETTINGS.hostname,
            "interval": SETTINGS.sample_interval,
        })

    @app.post("/api/pause")
    async def api_pause() -> JSONResponse:
        collector.pause()
        return JSONResponse({"paused": True})

    @app.post("/api/resume")
    async def api_resume() -> JSONResponse:
        collector.resume()
        return JSONResponse({"paused": False})

    # --- reguły normalizacji ------------------------------------------------
    def _rule_to_json(r) -> dict:
        return {
            "id": r.id, "enabled": r.enabled, "priority": r.priority,
            "app_match": r.app_match, "match_type": r.match_type,
            "pattern": r.pattern, "entity_tmpl": r.entity_tmpl,
            "category": r.category, "project_mode": r.project_mode,
            "project": r.project, "project_group": r.project_group,
            "is_decl": r.is_decl, "detail_group": r.detail_group,
            "note": r.note, "builtin": r.builtin,
        }

    def _rule_from_json(d: dict):
        from .rules import Rule
        allowed = {"pattern", "enabled", "priority", "app_match", "match_type",
                   "entity_tmpl", "category", "project_mode", "project",
                   "project_group", "is_decl", "detail_group", "note"}
        return Rule(**{k: v for k, v in d.items() if k in allowed})

    @app.get("/api/rules")
    async def api_rules_list() -> JSONResponse:
        from .rules import BUILTIN_RULES
        user = [_rule_to_json(r) for r in storage.list_rules()]
        builtin = [_rule_to_json(r) for r in BUILTIN_RULES]
        return JSONResponse({"rules": user, "builtin": builtin})

    @app.post("/api/rules")
    async def api_rules_create(request: Request) -> JSONResponse:
        body = await request.json()
        rule = _rule_from_json(body)
        if not rule.pattern and rule.match_type != "contains":
            raise HTTPException(400, "Pusty pattern")
        rid = storage.add_rule(rule)
        return JSONResponse({"id": rid, "ok": True})

    @app.put("/api/rules/{rule_id}")
    async def api_rules_update(rule_id: int, request: Request) -> JSONResponse:
        body = await request.json()
        ok = storage.update_rule(rule_id, **body)
        if not ok:
            raise HTTPException(404, "Reguła nie istnieje")
        return JSONResponse({"ok": True})

    @app.delete("/api/rules/{rule_id}")
    async def api_rules_delete(rule_id: int) -> JSONResponse:
        ok = storage.delete_rule(rule_id)
        if not ok:
            raise HTTPException(404, "Reguła nie istnieje")
        return JSONResponse({"ok": True})

    @app.post("/api/rules/test")
    async def api_rules_test(request: Request) -> JSONResponse:
        """Podgląd na żywo: co da draft reguły dla danego tytułu + czy deklaruje projekt."""
        from .rules import Rule
        body = await request.json()
        draft = _rule_from_json(body.get("rule") or {})
        app_name = body.get("app") or ""
        title = body.get("title") or ""
        tab = body.get("tab")
        out = {"app_match": draft.app_matches(app_name.lower())}
        if not out["app_match"]:
            return JSONResponse(out)
        from .rules import RuleEngine
        engine = RuleEngine([draft])
        norm = engine.apply(app_name, title, tab)
        decl = engine.detect_project(app_name, title, tab)
        out["matched"] = norm is not None
        if norm:
            out["entity"] = norm.entity
            out["category"] = norm.category
            out["project"] = norm.project
            out["detail"] = norm.detail
        out["declares_project"] = decl
        return JSONResponse(out)

    @app.post("/api/rules/test-coverage")
    async def api_rules_coverage(request: Request) -> JSONResponse:
        """Pokrycie draftu reguły na wszystkich tytułach wybranego dnia.

        Zwraca ile unikalnych tytułów pasuje + przykłady matchy i nie-matchy.
        """
        from .rules import RuleEngine
        body = await request.json()
        draft = _rule_from_json(body.get("rule") or {})
        date = body.get("date")
        if date:
            try:
                d = _dt.datetime.fromisoformat(date).astimezone()
            except ValueError:
                raise HTTPException(400, "Niepoprawna data")
        else:
            d = _dt.datetime.now().astimezone()

        start_utc, end_utc = storage._day_bounds_utc(d)
        s_iso = start_utc.isoformat(timespec="seconds")
        e_iso = end_utc.isoformat(timespec="seconds")
        with storage._cursor() as (_, cur):
            rows = cur.execute(
                """SELECT DISTINCT app, title, tab FROM events
                   WHERE ts_end > ? AND ts_start < ? AND idle = 0""",
                (s_iso, e_iso),
            ).fetchall()

        engine = RuleEngine([draft])
        matched, missed = [], []
        for r in rows:
            raw = (r["tab"] or r["title"] or "").strip()
            if not raw:
                continue
            if not draft.app_matches((r["app"] or "").lower()):
                continue
            norm = engine.apply(r["app"], r["title"], r["tab"])
            decl = engine.detect_project(r["app"], r["title"], r["tab"])
            if norm or decl:
                matched.append({
                    "app": r["app"], "title": raw,
                    "entity": norm.entity if norm else None,
                    "project": (norm.project if norm else None) or decl,
                })
            else:
                missed.append({"app": r["app"], "title": raw})
        return JSONResponse({
            "matched_count": len(matched),
            "missed_count": len(missed),
            "matched": matched[:15],
            "missed": missed[:15],
        })

    @app.post("/api/rules/suggest")
    async def api_rules_suggest(request: Request) -> JSONResponse:
        """Asystent: analizuje tytuły aplikacji i proponuje kandydatów na reguły.

        body: {app, days? (domyślnie 30), date? (koniec zakresu, domyślnie dziś)}
        """
        from .rules import suggest_for_app
        body = await request.json()
        app_name = (body.get("app") or "").strip().lower()
        if not app_name:
            raise HTTPException(400, "Podaj aplikację")
        days = int(body.get("days") or 30)
        days = max(1, min(days, 365))
        date = body.get("date")
        if date:
            try:
                end_d = _dt.datetime.fromisoformat(date).astimezone()
            except ValueError:
                raise HTTPException(400, "Niepoprawna data")
        else:
            end_d = _dt.datetime.now().astimezone()
        start_d = end_d - _dt.timedelta(days=days)

        with storage._cursor() as (_, cur):
            rows = cur.execute(
                """SELECT DISTINCT title, tab FROM events
                   WHERE lower(app) LIKE ? AND ts_start >= ? AND ts_start < ?
                     AND idle = 0""",
                (f"%{app_name}%", start_d.isoformat(timespec="seconds"),
                 end_d.isoformat(timespec="seconds")),
            ).fetchall()
        titles = [(r["tab"] or r["title"] or "").strip() for r in rows]
        titles = [t for t in titles if t]
        return JSONResponse({
            "app": app_name,
            "days": days,
            "titles_count": len(set(titles)),
            "candidates": suggest_for_app(titles),
        })

    # --- zarządzanie zbiorcze regułami (backup/clear/import) -----------------
    @app.get("/api/rules/export")
    async def api_rules_export() -> JSONResponse:
        """Eksport wszystkich reguł użytkownika do JSON (backup)."""
        rules = [_rule_to_json(r) for r in storage.list_rules()]
        for r in rules:
            r.pop("id", None)
        return JSONResponse(
            {"version": 1, "rules": rules},
            headers={"Content-Disposition": 'attachment; filename="timerecord-rules.json"'},
        )

    @app.post("/api/rules/clear")
    async def api_rules_clear() -> JSONResponse:
        """Wyzeruj reguły użytkownika — wcześniej archiwizuje do rules_archive_*.json."""
        import json as _json
        rules = [_rule_to_json(r) for r in storage.list_rules()]
        for r in rules:
            r.pop("id", None)
        archived_to = None
        if rules:
            from .config import DATA_DIR
            ts = _dt.datetime.now().strftime("%Y%m%d_%H%M%S")
            path = DATA_DIR / f"rules_archive_{ts}.json"
            path.write_text(_json.dumps({"version": 1, "rules": rules},
                                        ensure_ascii=False, indent=2),
                            encoding="utf-8")
            archived_to = str(path)
            for r in storage.list_rules():
                storage.delete_rule(r.id)
        return JSONResponse({"ok": True, "removed": len(rules),
                             "archived_to": archived_to})

    @app.post("/api/rules/import")
    async def api_rules_import(request: Request) -> JSONResponse:
        """Import reguł z JSON. mode: 'merge' (dodaj) | 'replace' (wyczyść+dodaj)."""
        body = await request.json()
        rules = body.get("rules") or []
        if not isinstance(rules, list):
            raise HTTPException(400, "Oczekiwano listy reguł")
        if body.get("mode") == "replace":
            for r in storage.list_rules():
                storage.delete_rule(r.id)
        n = storage.import_rules(rules)
        return JSONResponse({"ok": True, "imported": n})

    @app.get("/api/suggestions")
    async def api_suggestions(date: Optional[str] = None, limit: int = 30) -> JSONResponse:
        """Top nieklasyfikowane encje — kandydaci na reguły użytkownika."""
        if date:
            try:
                d = _dt.datetime.fromisoformat(date).astimezone()
            except ValueError:
                raise HTTPException(400, "Niepoprawna data")
        else:
            d = None
        rows = storage.unclassified_for_day(day=d, limit=min(max(1, limit), 100))
        return JSONResponse({"day": (d or _dt.datetime.now().astimezone()).date().isoformat(),
                             "suggestions": rows})

    def _enrich(summary: dict) -> dict:
        summary = dict(summary)
        summary["total_human"] = _fmt_duration(summary.get("total_seconds") or 0)
        summary["idle_human"] = _fmt_duration(summary.get("idle_seconds") or 0)
        for a in summary.get("apps", []):
            a["dur_human"] = _fmt_duration(a.get("dur_seconds") or 0)
        return summary

    return app
